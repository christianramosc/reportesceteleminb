# -*- coding: utf-8 -*-
"""Pruebas del modelo de datos (esquema estrella) y su exportación a Excel."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import construir_bitacora
from herramientas import modelo_datos as md
from herramientas.modelo_datos_excel import exportar_excel
from herramientas.herramienta_modelo_datos import ejecutar


def _dos_meses():
    """Junio y julio. El folio CCK0000 se aprueba en junio y se financia en
    julio; el cliente de CCK0001 vuelve en julio con otro folio."""
    junio = construir_bitacora(estatus=["APROBADO", "RECHAZADO", "FINANCIADO"], n=6, semilla=1)
    julio = construir_bitacora(estatus=["FINANCIADOS", "CONTRAPROPUESTA"], n=4, semilla=2)
    julio["Folio CCK"] = ["CCK0000", "CCK9001", "CCK9002", "CCK9003"]
    julio.loc[1, ["Nombre Completo", "Fecha de Nacimiento"]] = [
        junio.loc[1, "Nombre Completo"].upper(), junio.loc[1, "Fecha de Nacimiento"]]
    return [("JUNIO 2026", "junio.xlsx", junio), ("JULIO 2026", "julio.xlsx", julio)]


@pytest.fixture
def modelo():
    return md.construir_modelo(_dos_meses())


def test_cada_folio_cuenta_una_vez(modelo):
    h = modelo["hechos_solicitud"]
    assert len(modelo["hechos_estatus_mes"]) == 10
    assert len(h) == 9
    assert h.folio.is_unique


def test_folio_que_cruza_toma_ultimo_estatus_y_mes_de_captura(modelo):
    f = modelo["hechos_solicitud"].set_index("folio").loc["CCK0000"]
    assert f.estatus_captura == "APROBADO"
    assert f.estatus_final == "FINANCIADO"     # "FINANCIADOS" normalizado
    assert f.id_mes_captura == 202606 and f.id_mes_ultimo == 202607
    assert f.cruzo_mes == 1 and f.financiado == 1


def test_foto_mensual_marca_el_cambio(modelo):
    e = modelo["hechos_estatus_mes"]
    fila = e[(e.folio == "CCK0000") & (e.id_mes == 202607)].iloc[0]
    assert fila.estatus_anterior == "APROBADO" and fila.cambio_estatus == 1
    assert fila.es_ultimo == 1 and fila.es_captura == 0
    # filtrar es_ultimo = 1 reproduce hechos_solicitud
    assert e.es_ultimo.sum() == len(modelo["hechos_solicitud"])


def test_sin_datos_de_contacto(modelo):
    for nombre, tabla in modelo.items():
        for col in ("Nombre Completo", "Teléfono Móvil", "Correo Electrónico",
                    "Fecha de Nacimiento", "_nombre", "_nacimiento"):
            assert col not in tabla.columns, (nombre, col)


def test_recaptura_del_mismo_cliente(modelo):
    h = modelo["hechos_solicitud"].set_index("folio")
    assert h.loc["CCK9001", "es_recaptura"] == 1
    assert h.loc["CCK9001", "id_cliente"] == h.loc["CCK0001", "id_cliente"]
    assert h.loc["CCK0001", "es_recaptura"] == 0


def test_llaves_de_las_relaciones_existen(modelo):
    for tabla, col, destino, col_destino, _ in md.RELACIONES:
        hijos = set(modelo[tabla][col].dropna())
        padres = set(modelo[destino][col_destino])
        assert hijos <= padres, (tabla, col, hijos - padres)
        assert modelo[destino][col_destino].is_unique


def test_periodo_desde_hoja_o_archivo():
    assert md.interpretar_periodo("SEPTIEMBRE 2026")[:2] == (2026, 9)
    assert md.interpretar_periodo("Hoja1", "AGOSTO_2026_FINAL.xlsx")[:2] == (2026, 8)
    anio, mes, supuesto = md.interpretar_periodo("JULIO", anio_defecto=2025)
    assert (anio, mes, supuesto) == (2025, 7, True)


def test_mes_irreconocible_da_error_claro():
    with pytest.raises(ValueError, match="No se reconoce el mes"):
        md.construir_modelo([("Hoja1", "datos.xlsx", construir_bitacora(n=3))])


def test_folio_repetido_en_el_mes_se_avisa():
    df = construir_bitacora(n=4)
    df.loc[3, "Folio CCK"] = df.loc[0, "Folio CCK"]
    m = md.construir_modelo([("JUNIO 2026", "junio.xlsx", df)])
    assert len(m["hechos_solicitud"]) == 3
    assert "Folio repetido en el mes" in set(m["calidad_datos"].tipo)


def test_vehiculo_se_descompone():
    assert md.descomponer_vehiculo("NEW ZS 1.5T COM TURBO EXCITE AT") == ("ZS", "AT", "NO")
    assert md.descomponer_vehiculo("MG3 1.5L HEV ELEGANCE")[2] == "SI"
    assert md.descomponer_vehiculo("MG5 STYLE")[1] == md.SIN_DATO


def test_excel_trae_tablas_con_nombre(modelo, tmp_path):
    ruta = exportar_excel(modelo, str(tmp_path / "modelo.xlsx"), ["junio.xlsx", "julio.xlsx"])
    wb = load_workbook(ruta)
    tablas = {n for ws in wb for n in ws.tables.keys()}
    assert {"hechos_solicitud", "hechos_estatus_mes", "dim_mes", "dim_vendedor",
            "dim_vehiculo", "dim_estatus", "dim_cliente"} <= tablas
    assert wb.sheetnames[0] == "Léeme"
    assert wb["Cuadre"]["C5"].value.startswith("=COUNTIFS(")


def test_herramienta_de_la_app(tmp_path):
    rutas = []
    for etiqueta, archivo, df in _dos_meses():
        ruta = tmp_path / f"0_{archivo}"
        with pd.ExcelWriter(ruta) as w:
            df.to_excel(w, sheet_name=etiqueta, index=False)
        rutas.append(str(ruta))
    salida = ejecutar(rutas, carpeta_salida=str(tmp_path / "salida"))
    assert len(salida) == 1 and salida[0].endswith("Modelo_Datos_Bitacora_Jun26-Jul26.xlsx")
