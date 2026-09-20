# -*- coding: utf-8 -*-
"""
Pruebas de generación de reportes de punta a punta.

Son las más lentas (generan PDF y Word de verdad), pero son las que valen:
verifican lo que el usuario final recibe.
"""

import contextlib
import glob
import io
import os
import shutil

import pytest

from herramientas import REGISTRO
from herramientas import avance_preliminar_original as avance


def _herramienta(identificador):
    return next(h for h in REGISTRO if h.id == identificador)


def _limpiar_temporales_como_la_app():
    """Reproduce la limpieza que hace streamlit_app.py después de cada
    corrida. Es indispensable para que la prueba refleje la realidad."""
    for carpeta in glob.glob("graficas_temp_*"):
        shutil.rmtree(carpeta, ignore_errors=True)


def _contar_imagenes(ruta_pdf):
    """Cuenta los PNG incrustados, sin depender de herramientas externas."""
    with open(ruta_pdf, "rb") as f:
        return f.read().count(b"/Subtype /Image")


IDS_REPORTES = ["avance_preliminar", "resumen_mensual", "comparativo_mensual"]


def _ejecutar(identificador, bitacora, destino):
    herramienta = _herramienta(identificador)
    entradas = [bitacora, bitacora] if identificador == "comparativo_mensual" else [bitacora]
    opciones = {"nombres_meses": ["MES A", "MES B"]} if identificador == "comparativo_mensual" else {}
    with contextlib.redirect_stdout(io.StringIO()):
        return herramienta.ejecutar(entradas, carpeta_salida=destino, **opciones)


@pytest.mark.parametrize("identificador", IDS_REPORTES)
def test_genera_pdf_y_word(identificador, bitacora_excel, tmp_path):
    destino = str(tmp_path / identificador)
    try:
        rutas = _ejecutar(identificador, bitacora_excel, destino)
    finally:
        _limpiar_temporales_como_la_app()

    extensiones = sorted(os.path.splitext(r)[1] for r in rutas)
    assert extensiones == [".docx", ".pdf", ".xlsx"], f"{identificador} -> {rutas}"
    for ruta in rutas:
        assert os.path.getsize(ruta) > 10_000, f"{ruta} salió sospechosamente chico"


@pytest.mark.parametrize("identificador", IDS_REPORTES)
def test_el_libro_lleva_datos_y_formulas(identificador, bitacora_excel, tmp_path):
    """El Excel debe traer la hoja Datos y tablas con fórmulas vivas.

    Se exige que las celdas sean fórmulas y no números pegados: si alguien
    "arregla" el exportador escribiendo valores calculados en Python, el
    libro seguiría abriéndose bien pero dejaría de recalcular, que es justo
    lo que lo hace útil.
    """
    from openpyxl import load_workbook

    destino = str(tmp_path / identificador)
    try:
        rutas = _ejecutar(identificador, bitacora_excel, destino)
    finally:
        _limpiar_temporales_como_la_app()

    xlsx = next(r for r in rutas if r.endswith(".xlsx"))
    wb = load_workbook(xlsx)
    assert "Datos" in wb.sheetnames
    assert "Léeme" in wb.sheetnames
    assert wb["Datos"].max_row > 1, "la hoja Datos salió vacía"

    formulas = [
        celda.value
        for nombre in wb.sheetnames if nombre not in ("Datos", "Léeme")
        for fila in wb[nombre].iter_rows()
        for celda in fila
        if isinstance(celda.value, str) and celda.value.startswith("=")
    ]
    assert len(formulas) >= 20, f"{identificador}: solo {len(formulas)} fórmulas"
    assert any("COUNTIFS" in f for f in formulas)
    assert wb.calculation.fullCalcOnLoad, "sin esto Excel abre las celdas en blanco"


@pytest.mark.parametrize("identificador", IDS_REPORTES)
def test_el_pdf_lleva_graficas(identificador, bitacora_excel, tmp_path):
    """Un PDF sin gráficas se ve normal pero está mutilado: hay que exigir
    que las traiga, no solo que el archivo exista."""
    destino = str(tmp_path / identificador)
    try:
        rutas = _ejecutar(identificador, bitacora_excel, destino)
    finally:
        _limpiar_temporales_como_la_app()

    pdf = next(r for r in rutas if r.endswith(".pdf"))
    assert _contar_imagenes(pdf) >= 5


@pytest.mark.parametrize("identificador", IDS_REPORTES)
def test_dos_corridas_seguidas_en_el_mismo_proceso(identificador, bitacora_excel, tmp_path):
    """LA prueba de regresión de este proyecto.

    La app de Streamlit vive en un proceso de larga vida y borra los
    temporales entre corridas. Una vez, la carpeta de gráficas se creaba
    solo al importar el módulo: la primera corrida funcionaba y la segunda
    tronaba con 'No such file or directory' (o peor, el comparativo entregaba
    un PDF sin una sola gráfica, sin avisar). Probar UNA vez no lo detecta.
    """
    imagenes = []
    for vuelta in (1, 2):
        destino = str(tmp_path / f"{identificador}_{vuelta}")
        try:
            rutas = _ejecutar(identificador, bitacora_excel, destino)
        finally:
            _limpiar_temporales_como_la_app()
        pdf = next(r for r in rutas if r.endswith(".pdf"))
        imagenes.append(_contar_imagenes(pdf))

    assert imagenes[1] >= 5, f"la 2a corrida perdió las gráficas: {imagenes}"
    assert imagenes[0] == imagenes[1], f"corridas distintas: {imagenes}"


def test_sigue_funcionando_con_los_estatus_originales(bitacora_clasica_excel, tmp_path):
    """Al ampliar el catálogo a 14 estatus, una bitácora vieja de 4 debe
    seguir generando su reporte igual."""
    try:
        rutas = _ejecutar("avance_preliminar", bitacora_clasica_excel, str(tmp_path))
    finally:
        _limpiar_temporales_como_la_app()
    assert len(rutas) == 3   # PDF, Word y el libro de análisis


def test_avisa_si_no_puede_generar_ninguna_grafica():
    """Antes se armaba un PDF completo sin una sola gráfica y sin error."""
    import pandas as pd
    df = pd.DataFrame({"Folio CCK": ["A", "B"], "Otro": [1, 2]})
    with contextlib.redirect_stdout(io.StringIO()):
        limpio = avance.limpiar_datos(df)
        with pytest.raises(RuntimeError, match="ninguna gráfica"):
            avance.generar_graficas_para_pdf(limpio)
    _limpiar_temporales_como_la_app()
